from flask import Flask, render_template, request, send_file
import os
import docx
from fpdf import FPDF

app = Flask(__name__)

# Function to extract Wi-Fi details from .docx
def extract_wifi_details(docx_path):
    doc = docx.Document(docx_path)
    wifi_details = {"username": "", "password": ""}
    
    # Loop through paragraphs in the docx to find relevant details
    for para in doc.paragraphs:
        text = para.text.strip()
        if 'username' in text.lower():
            wifi_details['username'] = text.split(':')[-1].strip()
        elif 'password' in text.lower():
            wifi_details['password'] = text.split(':')[-1].strip()
    
    return wifi_details

# Function to create a PDF with Wi-Fi details
def save_to_pdf(pdf_path, wifi_details):
    pdf = FPDF()
    pdf.add_page()

    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Wi-Fi Details", ln=True, align='C')
    pdf.ln(10)
    pdf.cell(200, 10, txt=f"Username: {wifi_details['username']}", ln=True)
    pdf.cell(200, 10, txt=f"Password: {wifi_details['password']}", ln=True)

    pdf.output(pdf_path)

@app.route('/', methods=['GET'])
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return "No file part"
    
    file = request.files['file']

    if file.filename == '':
        return "No selected file"
    
    if file and file.filename.endswith('.docx'):
        # Save the uploaded file
        filepath = os.path.join('uploads', file.filename)
        file.save(filepath)

        # Extract Wi-Fi details and generate PDF
        wifi_details = extract_wifi_details(filepath)
        pdf_filename = f"WiFi_{file.filename.split('.')[0]}.pdf"
        pdf_path = os.path.join('uploads', pdf_filename)
        save_to_pdf(pdf_path, wifi_details)

        # Serve the generated PDF back to the user
        return send_file(pdf_path, as_attachment=True)

if __name__ == '__main__':
    if not os.path.exists('uploads'):
        os.makedirs('uploads')
    app.run(debug=True)

@app.route('/upload', methods=['POST'])
def upload_file():
    if request.method == 'POST':
        file = request.files['file']
        if file:
            filepath = os.path.join(upload_file, file.filename)
            file.save(filepath)
            return "File uploaded successfully"

from fpdf import FPDF

# Function to create the PDF with Wi-Fi details
def create_wifi_pdf(pdf_path, wifi_data):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    for entry in wifi_data:
        pdf.add_page()
        
        # Set Title
        pdf.set_font("Arial", "B", 16)
        pdf.cell(200, 10, txt="Wi-Fi Name: Network", ln=True, align="C")
        pdf.cell(200, 10, txt=f"Password: {entry['network_password']}", ln=True, align="C")
        pdf.ln(10)

        # Add Login Credentials Section
        pdf.set_font("Arial", "B", 12)
        pdf.cell(200, 10, txt="Login Credentials", ln=True, align="L")
        
        pdf.set_font("Arial", "", 12)
        pdf.cell(200, 10, txt=f"Username: {entry['username']}", ln=True)
        pdf.cell(200, 10, txt=f"Password: {entry['password']}", ln=True)
        
        pdf.ln(20)

        # Add SCAN ME section
        pdf.set_font("Arial", "B", 16)
        pdf.cell(200, 10, txt="SCAN ME", ln=True, align="C")
        pdf.cell(200, 10, txt="Welcome to ZINGSA", ln=True, align="C")
    
    # Output to the PDF file
    pdf.output(pdf_path)

# Sample Wi-Fi data based on the example provided
wifi_data = [
    {"username": "guest-00193", "password": "0KWaq2dW", "network_password": "T!@_zing2a"},
    {"username": "guest-00194", "password": "8mfkRlSL", "network_password": "T!@_zing2a"},
    {"username": "guest-00195", "password": "Zckfa7cE", "network_password": "T!@_zing2a"},
    # Add more entries as needed
]

# Path to save the PDF
pdf_path = r"C:\Users\user\Desktop\WiFi09_Generated.pdf"

# Create the PDF
create_wifi_pdf(pdf_path, wifi_data)

print(f"PDF generated and saved to {pdf_path}")
